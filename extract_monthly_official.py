"""Normalize monthly NSO 2026 trade and industrial production workbooks.

Inputs: data/nso/raw/monthly/E01-2026-7.xls, E02-2026-7.xls, IIP-Eng9.xlsx
Requires: xlrd>=2.0.1, openpyxl>=3.1
"""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

import openpyxl
import xlrd
from data_paths import data_path

ROOT = Path(__file__).resolve().parent
SOURCES = ROOT / "data" / "nso" / "raw" / "monthly"


def write_csv(name: str, rows: list[list[object]]) -> None:
    with data_path(name, for_write=True).open("w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)


def parse_trade(path: Path) -> dict[int, float]:
    book = xlrd.open_workbook(str(path), on_demand=True)
    sheet = book.sheet_by_index(0)
    month_cols = {}
    for col, value in enumerate(sheet.row_values(2)):
        month = str(value).strip().lower()
        if month in {"january", "february", "march", "april", "may", "june", "july", "august", "september"}:
            month_cols[col] = col + 1  # NSO layout: quantity, then value (1000 USD).
    total_row = next((i for i in range(3, sheet.nrows) if str(sheet.cell_value(i, 0)).strip().lower() == "total"), None)
    if total_row is None or len(month_cols) != 9:
        raise ValueError(f"Unexpected NSO trade layout in {path.name}")
    values = {}
    for col, value_col in month_cols.items():
        month_name = str(sheet.cell_value(2, col)).strip().lower()
        month = ["january", "february", "march", "april", "may", "june", "july", "august", "september"].index(month_name) + 1
        amount_1000_usd = float(sheet.cell_value(total_row, value_col))
        values[month] = amount_1000_usd * 1000
    book.release_resources()
    return values


exports = parse_trade(SOURCES / "E01-2026-7.xls")
imports = parse_trade(SOURCES / "E02-2026-7.xls")
if set(exports) != set(range(1, 10)) or set(imports) != set(range(1, 10)):
    raise ValueError("Monthly trade coverage must be Jan-Sep 2026")

trade_rows: list[list[object]] = [[
    "thoi_gian", "kim_ngach_xuat_khau_usd", "kim_ngach_nhap_khau_usd",
    "can_can_thuong_mai_usd", "tan_suat", "trang_thai", "don_vi_nguon",
]]
for month in range(1, 10):
    export, import_ = exports[month], imports[month]
    trade_rows.append([
        date(2026, month, 1).isoformat(), round(export, 3), round(import_, 3),
        round(export - import_, 3), "thang", "so_bo", "1000_USD_converted_to_USD",
    ])
write_csv("du_lieu_xuat_nhap_khau_thang_2026.csv", trade_rows)


book = openpyxl.load_workbook(SOURCES / "IIP-Eng9.xlsx", read_only=True, data_only=True)
sheet = book["IIP (C4)"]
if sheet.cell(6, 2).value != "Whole industry":
    raise ValueError("Expected Whole industry row in NSO IIP workbook")
iip_months = [sheet.cell(6, col).value for col in range(3, 12)]
if len(iip_months) != 9 or any(value is None for value in iip_months):
    raise ValueError("Expected nine monthly IIP observations Jan-Sep 2026")
mom_index = sheet.cell(6, 12).value  # Sep 2026 over Aug 2026, in percent.
yoy_index = sheet.cell(6, 13).value  # Sep 2026 over Sep 2025, in percent.
nine_month_index = sheet.cell(6, 14).value  # Jan-Sep 2026 over Jan-Sep 2025.
iip_rows: list[list[object]] = [[
    "thoi_gian", "chi_so_iip_toan_nganh_2015_100", "tang_truong_mom_pct",
    "tang_truong_yoy_pct", "tan_suat", "trang_thai", "don_vi_nguon",
]]
for month, value in enumerate(iip_months, start=1):
    iip_rows.append([
        date(2026, month, 1).isoformat(), float(value),
        round(float(mom_index) - 100, 2) if month == 9 else "",
        round(float(yoy_index) - 100, 2) if month == 9 else "",
        "thang", "so_bo", "chi_so_2015_binh_quan_thang_100",
    ])
write_csv("du_lieu_san_xuat_iip_thang_2026.csv", iip_rows)

write_csv("du_lieu_san_xuat_iip_so_sanh_2026.csv", [
    ["chi_tieu", "gia_tri_pct", "ky_so_sanh", "trang_thai", "ghi_chu"],
    ["toan_nganh_iip_tang_truong", round(float(mom_index) - 100, 2), "2026-09_so_voi_2026-08", "so_bo", "Chi so NSO cong bo la 103.26; chuyen ve tang truong 3.26%"],
    ["toan_nganh_iip_tang_truong", round(float(yoy_index) - 100, 2), "2026-09_so_voi_2025-09", "so_bo", "Chi so NSO cong bo la 116.65; chuyen ve tang truong 16.65%"],
    ["toan_nganh_iip_tang_truong", round(float(nine_month_index) - 100, 2), "2026-01_den_2026-09_so_voi_cung_ky_2025", "so_bo", "Chi so NSO cong bo la 112.33; chuyen ve tang truong 12.33%"],
])
book.close()

# Compatibility table matching the project's combined IIP/retail/trade schema.
with data_path("du_lieu_ban_le_tieu_dung_thang_2026.csv").open(encoding="utf-8-sig", newline="") as f:
    retail_by_date = {r["thoi_gian"]: r for r in csv.DictReader(f)}
iip_by_date = {row[0]: row for row in iip_rows[1:]}
compat_rows: list[list[object]] = [[
    "thoi_gian", "chi_so_iip", "tong_muc_ban_le_vnd", "kim_ngach_xuat_khau_usd",
    "kim_ngach_nhap_khau_usd", "can_can_thuong_mai_usd", "tang_truong_danh_nghia_ban_le_yoy_pct",
    "tang_truong_thuc_ban_le_yoy_luy_ke_pct", "trang_thai_iip", "trang_thai_thuong_mai", "trang_thai_ban_le",
]]
for row in trade_rows[1:]:
    stamp = str(row[0])
    iip_row = iip_by_date[stamp]
    retail_row = retail_by_date.get(stamp, {})
    compat_rows.append([
        stamp, iip_row[1], retail_row.get("tong_muc_ban_le_vnd", ""), row[1], row[2], row[3],
        retail_row.get("tang_truong_danh_nghia_yoy_pct", ""),
        retail_row.get("tang_truong_thuc_yoy_luy_ke_pct", ""),
        "so_bo", row[5], retail_row.get("trang_thai", ""),
    ])
write_csv("du_lieu_iip_ban_le_xuat_nhap_khau.csv", compat_rows)

print(f"Trade rows: {len(trade_rows)-1}; exports 9m={sum(exports.values()):,.0f} USD; imports 9m={sum(imports.values()):,.0f} USD")
print(f"IIP rows: {len(iip_rows)-1}; Sep index={iip_months[-1]}, Sep MoM={float(mom_index)-100:.2f}%, YoY={float(yoy_index)-100:.2f}%")
