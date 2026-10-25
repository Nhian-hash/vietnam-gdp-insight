"""Normalize and analyze additional official GDP growth-driver datasets."""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from data_paths import data_path

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "data" / "nso" / "raw" / "expanded"
NSO_ACCOUNTS = "https://pxweb.nso.gov.vn/pxweb/en/E%2003.%20National%20Accounts%20and%20State%20budget/E%2003.%20National%20Accounts%20and%20State%20budget__Px_E_03/?tablelist=true"
NSO_LABOR = "https://pxweb.nso.gov.vn/pxweb/en/E%2002.%20Population%20and%20Employment/E%2002.%20Population%20and%20Employment__Px_E_02/?tablelist=true"
NSO_INVESTMENT = "https://pxweb.nso.gov.vn/pxweb/en/E%2004.%20Investment%20and%20construction/E%2004.%20Investment%20and%20construction__Px_E_04/?tablelist=true"
WORLD_BANK_COMMODITIES = "https://www.worldbank.org/en/research/commodity-markets"


def read_px(name: str) -> pd.DataFrame:
    path = SRC / f"{name}.csv"
    return pd.read_csv(path, skiprows=2, dtype=str, encoding="utf-8-sig").dropna(how="all")


def num(value):
    if pd.isna(value):
        return np.nan
    text = str(value).strip().replace(",", "")
    if text in {"", "..", "…", "-", "—"}:
        return np.nan
    try:
        return float(text)
    except ValueError:
        return np.nan


def year_col(col):
    match = re.search(r"(?:19|20)\d{2}", str(col))
    return int(match.group()) if match else None


def status_for(label, year):
    text = str(label).lower()
    if "est" in text:
        return "uoc_tinh"
    if "prel" in text or "preliminary" in text:
        return "so_bo"
    return "theo_bang_nguon"


def save(df: pd.DataFrame, name: str):
    df.to_csv(data_path(name, for_write=True), index=False, encoding="utf-8-sig", lineterminator="\n")


def melt_table(df, id_cols, label_col, value_type_func):
    records = []
    year_columns = [(col, year_col(col)) for col in df.columns if year_col(col)]
    for _, row in df.iterrows():
        label = row[label_col]
        for col, year in year_columns:
            value = num(row[col])
            if not np.isfinite(value):
                continue
            extra = {key: row[key] for key in id_cols}
            extra.update({"nam": year, "chi_tieu": label, "gia_tri": value,
                          "loai_chi_tieu": value_type_func(row, col),
                          "trang_thai": status_for(col, year)})
            records.append(extra)
    return pd.DataFrame(records)


def build_gdp_expenditure():
    raw = read_px("gdp_constant_expenditure")
    records = []
    for _, row in raw.iterrows():
        measure = str(row.iloc[0])
        category = str(row.iloc[1])
        is_index = measure.lower().startswith("index")
        for col in raw.columns[2:]:
            year = year_col(col)
            value = num(row[col])
            if year is None or not np.isfinite(value):
                continue
            records.append({
                "nam": year, "chi_tieu": category,
                "gia_tri_vnd_ty_gia_2010": value if not is_index else np.nan,
                "chi_so_nam_truoc_100": value if is_index else np.nan,
                "tang_truong_yoy_pct": value - 100 if is_index else np.nan,
                "don_vi": "tỷ VND giá cố định 2010" if not is_index else "% so với năm trước, chỉ số - 100",
                "loai_chi_tieu": "tang_truong_chinh_thuc" if is_index else "muc_gia_thuc",
                "trang_thai": status_for(col, year),
            })
    result = pd.DataFrame(records)
    save(result, "du_lieu_gdp_theo_chi_tieu.csv")

    leaves = [
        "Gross capital formation- Gross  fixed capital formation",
        "Gross capital formation- Changes in inventories",
        "Final consumption- State",
        "Final consumption- Household",
        "Balance of Trade (goods & services)",
        "Statistical discrepancy",
    ]
    levels = result[result["loai_chi_tieu"].eq("muc_gia_thuc")].copy()
    total = levels[levels["chi_tieu"].eq("TOTAL")].set_index("nam")["gia_tri_vnd_ty_gia_2010"]
    pivot = levels[levels["chi_tieu"].isin(leaves)].pivot(index="nam", columns="chi_tieu", values="gia_tri_vnd_ty_gia_2010")
    official_growth = result[(result["loai_chi_tieu"].eq("tang_truong_chinh_thuc")) & result["chi_tieu"].eq("TOTAL")].set_index("nam")["tang_truong_yoy_pct"]
    contrib_rows = []
    for year in sorted(set(pivot.index) & set(total.index)):
        if year - 1 not in total.index:
            continue
        previous = total.loc[year - 1]
        if not np.isfinite(previous) or previous == 0:
            continue
        record = {"nam": int(year), "tang_truong_gdp_thuc_chi_tieu_pct": num(official_growth.get(year, np.nan)),
                  "gdp_thuc_vnd_ty_gia_2010": total.loc[year], "trang_thai": "uoc_tinh" if year == 2025 else ("so_bo" if year == 2024 else "theo_bang_nguon")}
        if year == 2010:
            # 2010 is the chain's change of reference/base year; do not compare 2009 and 2010 level series.
            for category in leaves:
                stem = {
                    "Gross capital formation- Gross  fixed capital formation": "von_co_dinh",
                    "Gross capital formation- Changes in inventories": "ton_kho",
                    "Final consumption- State": "tieu_dung_nha_nuoc",
                    "Final consumption- Household": "tieu_dung_ho_gia_dinh",
                    "Balance of Trade (goods & services)": "can_can_hang_hoa_dich_vu",
                    "Statistical discrepancy": "chenh_lech_thong_ke",
                }[category]
                record[f"muc_{stem}_vnd_ty_gia_2010"] = np.nan
                record[f"dong_gop_{stem}_diem_pct"] = np.nan
            record["tong_dong_gop_phan_ra_diem_pct"] = np.nan
            record["ghi_chu"] = "Không tính đóng góp 2010/so với 2009 vì chuyển năm gốc sang giá cố định 2010; giữ tốc độ tăng chính thức từ bảng NSO."
            contrib_rows.append(record)
            continue
        for category in leaves:
            stem = {
                "Gross capital formation- Gross  fixed capital formation": "von_co_dinh",
                "Gross capital formation- Changes in inventories": "ton_kho",
                "Final consumption- State": "tieu_dung_nha_nuoc",
                "Final consumption- Household": "tieu_dung_ho_gia_dinh",
                "Balance of Trade (goods & services)": "can_can_hang_hoa_dich_vu",
                "Statistical discrepancy": "chenh_lech_thong_ke",
            }[category]
            value = pivot.loc[year, category] if category in pivot.columns and year in pivot.index else np.nan
            previous_value = pivot.loc[year - 1, category] if category in pivot.columns and year - 1 in pivot.index else np.nan
            record[f"muc_{stem}_vnd_ty_gia_2010"] = value
            record[f"dong_gop_{stem}_diem_pct"] = (value - previous_value) / previous * 100 if np.isfinite(value) and np.isfinite(previous_value) else np.nan
        record["tong_dong_gop_phan_ra_diem_pct"] = sum(record.get(f"dong_gop_{key}_diem_pct", 0) for key in ["von_co_dinh", "ton_kho", "tieu_dung_nha_nuoc", "tieu_dung_ho_gia_dinh", "can_can_hang_hoa_dich_vu", "chenh_lech_thong_ke"] if np.isfinite(record.get(f"dong_gop_{key}_diem_pct", np.nan)))
        record["ghi_chu"] = ""
        contrib_rows.append(record)
    contrib = pd.DataFrame(contrib_rows)
    contrib["sai_lech_phan_ra_diem_pct"] = contrib["tong_dong_gop_phan_ra_diem_pct"] - contrib["tang_truong_gdp_thuc_chi_tieu_pct"]
    save(contrib, "du_lieu_dong_gop_gdp_theo_chi_tieu.csv")
    return result, contrib


def build_real_sectors():
    raw = read_px("gdp_constant_sector")
    records = []
    sectors = [
        ("Total (Bill. Dongs) Total", "Tổng GDP", "tong"),
        ("Total (Bill. Dongs) Agriculture, forestry and fishing", "Nông, lâm, thủy sản", "nong_nghiep"),
        ("Total (Bill. Dongs) Industry and construction", "Công nghiệp và xây dựng", "cong_nghiep_xay_dung"),
        ("Total (Bill. Dongs) Services", "Dịch vụ", "dich_vu"),
        ("Total (Bill. Dongs) Taxes less subsidies on products", "Thuế sản phẩm trừ trợ cấp", "thue_san_pham"),
    ]
    for _, row in raw.iterrows():
        year = year_col(row.iloc[0])
        if year is None:
            continue
        for level_col, label, key in sectors:
            growth_col = next((c for c in raw.columns if "Index (Previous year" in c and (label_map := key) and (
                (key == "tong" and c.endswith("Total")) or
                (key == "nong_nghiep" and "Agriculture" in c) or
                (key == "cong_nghiep_xay_dung" and "Industry and construction" in c) or
                (key == "dich_vu" and c.endswith("Services")) or
                (key == "thue_san_pham" and "Taxes less subsidies" in c)
            )), None)
            level = num(row.get(level_col, np.nan))
            growth_index = num(row.get(growth_col, np.nan)) if growth_col else np.nan
            records.append({"nam": year, "nganh": label, "ma_nganh": key,
                            "gdp_thuc_vnd_ty_gia_2010": level,
                            "chi_so_tang_truong_nam_truoc_100": growth_index,
                            "tang_truong_yoy_pct": growth_index - 100 if np.isfinite(growth_index) else np.nan,
                            "don_vi": "tỷ VND giá cố định 2010",
                            "trang_thai": status_for(row.iloc[0], year)})
    result = pd.DataFrame(records)
    save(result, "du_lieu_gdp_thuc_theo_nganh.csv")
    return result


def build_detailed_gdp_activities():
    """Normalize NSO value-added levels, current-price structure and real growth by activity."""
    current = read_px("gdp_current_activity")
    real = read_px("gdp_constant_growth_activity")
    activity_names = {
        "Agriculture, forestry and fishing": ("agriculture", "Nông, lâm nghiệp và thủy sản"),
        "Mining and quarrying": ("industry", "Khai khoáng"),
        "Manufacturing": ("industry", "Công nghiệp chế biến, chế tạo"),
        "Electricity, gas, stream and air conditioning supply": ("industry", "Sản xuất và phân phối điện, khí đốt, hơi nước và điều hòa không khí"),
        "Water supply, sewerage, waste management and remediation activities": ("industry", "Cung cấp nước; quản lý và xử lý rác thải, nước thải"),
        "Construction": ("industry", "Xây dựng"),
        "Wholesale and retail trade; repair of motor vehicles and motorcycles": ("services", "Bán buôn, bán lẻ; sửa chữa ô tô, mô tô, xe máy"),
        "Transportation and storage": ("services", "Vận tải, kho bãi"),
        "Accommodation and food service activities": ("services", "Dịch vụ lưu trú và ăn uống"),
        "Information and communication": ("services", "Thông tin và truyền thông"),
        "Financial, banking and insurance activities": ("services", "Tài chính, ngân hàng và bảo hiểm"),
        "Real estate activities": ("services", "Kinh doanh bất động sản"),
        "Professional, scientific and technical activities": ("services", "Chuyên môn, khoa học và công nghệ"),
        "Administrative and support service activities": ("services", "Hành chính và dịch vụ hỗ trợ"),
        "Activities of Communist Party, socio-political organizations; public administration and defence; compulsory security": ("services", "Đảng, tổ chức chính trị-xã hội; quản lý nhà nước, quốc phòng và an ninh; bảo đảm xã hội bắt buộc"),
        "Education and training": ("services", "Giáo dục và đào tạo"),
        "Human health and social work activities": ("services", "Y tế và hoạt động trợ giúp xã hội"),
        "Arts, entertainment and recreation": ("services", "Nghệ thuật, vui chơi và giải trí"),
        "Other service activities": ("services", "Hoạt động dịch vụ khác"),
        "Activities of households as employers; undifferentiated goods and services producing activities of households for own use": ("services", "Làm thuê trong hộ; sản xuất hàng hóa, dịch vụ tự tiêu dùng của hộ"),
        "Taxes less subsidies on products": ("taxes", "Thuế sản phẩm trừ trợ cấp"),
    }
    sector_names = {"agriculture": "Nông, lâm, thủy sản", "industry": "Công nghiệp và xây dựng",
                    "services": "Dịch vụ", "taxes": "Thuế sản phẩm trừ trợ cấp"}

    def extract(raw, measure, unit_type):
        subset = raw[raw.iloc[:, 0].astype(str).str.strip().eq(measure)]
        output = {}
        for _, row in subset.iterrows():
            # NSO repeats the tax line once above the activity block and once inside it.
            # Both are the same GDP component; canonicalize the label so it is counted once.
            item = str(row.iloc[1]).strip().rstrip(".")
            if item not in activity_names:
                continue
            for col in raw.columns[2:]:
                year = year_col(col)
                value = num(row[col])
                if year is not None and np.isfinite(value):
                    output[(item, year)] = (value, str(col))
        return output

    current_values = extract(current, "Value (Bill.dongs)", "current")
    current_shares = extract(current, "Structure (%)", "share")
    real_values = extract(real, "Value (Bill.dongs)", "real")
    growth_indices = extract(real, "Index (Previous year = 100) (%)", "growth")
    records = []
    for (activity, year), (nominal, nominal_col) in current_values.items():
        sector_key, activity_vi = activity_names[activity]
        real_value, real_col = real_values.get((activity, year), (np.nan, nominal_col))
        share, _ = current_shares.get((activity, year), (np.nan, nominal_col))
        index, growth_col = growth_indices.get((activity, year), (np.nan, nominal_col))
        growth = index - 100 if np.isfinite(index) else np.nan
        records.append({
            "nam": year, "khu_vuc_ma": sector_key, "khu_vuc": sector_names[sector_key],
            "hoat_dong_kinh_te": activity_vi, "hoat_dong_kinh_te_nguon": activity,
            "gdp_hien_hanh_ty_vnd": nominal, "ty_trong_gdp_hien_hanh_pct": share,
            "gdp_thuc_ty_vnd_theo_bang_nguon": real_value,
            "chi_so_gdp_thuc_nam_truoc_100": index, "tang_truong_gdp_thuc_yoy_pct": growth,
            "tang_truong_danh_nghia_yoy_pct": np.nan, "dong_gop_gdp_tang_truong_diem_pct": np.nan,
            "trang_thai": status_for(nominal_col, year),
            "ghi_chu": "Giá hiện hành và tỷ trọng là danh nghĩa; tăng trưởng thực lấy từ chỉ số giá cố định của NSO. Mức giá cố định trước/sau năm gốc 2010 không nối trực tiếp để tính tăng trưởng năm 2010.",
        })
    result = pd.DataFrame(records).sort_values(["khu_vuc_ma", "hoat_dong_kinh_te", "nam"]).reset_index(drop=True)
    result["tang_truong_danh_nghia_yoy_pct"] = result.groupby("hoat_dong_kinh_te_nguon")["gdp_hien_hanh_ty_vnd"].pct_change() * 100

    # The official fixed-price levels support exact contribution arithmetic within a stable base-price segment.
    total_levels = real_values.get(("TOTAL", 0))  # Intentionally absent; use the separately published total row below.
    total_real = {}
    for _, row in real[real.iloc[:, 1].astype(str).str.strip().eq("TOTAL") & real.iloc[:, 0].astype(str).str.startswith("Value")].iterrows():
        for col in real.columns[2:]:
            year = year_col(col); value = num(row[col])
            if year is not None and np.isfinite(value): total_real[year] = value
    del total_levels
    for idx, row in result.iterrows():
        year = int(row["nam"])
        prior = year - 1
        current_value = row["gdp_thuc_ty_vnd_theo_bang_nguon"]
        prior_value = real_values.get((row["hoat_dong_kinh_te_nguon"], prior), (np.nan, ""))[0]
        denominator = total_real.get(prior, np.nan)
        # 2010 contains a national-accounts base/reference-year break in the level series.
        if year != 2010 and np.isfinite(current_value) and np.isfinite(prior_value) and np.isfinite(denominator) and denominator != 0:
            result.at[idx, "dong_gop_gdp_tang_truong_diem_pct"] = (current_value - prior_value) / denominator * 100
    result["bien_doi_ty_trong_5_nam_diem_pct"] = result.groupby("hoat_dong_kinh_te_nguon")["ty_trong_gdp_hien_hanh_pct"].transform(lambda s: s - s.shift(5))
    result["bien_doi_ty_trong_10_nam_diem_pct"] = result.groupby("hoat_dong_kinh_te_nguon")["ty_trong_gdp_hien_hanh_pct"].transform(lambda s: s - s.shift(10))
    result["don_vi_gdp"] = "tỷ VND"
    save(result, "du_lieu_gdp_nganh_chi_tiet.csv")
    return result


def build_gdp_activity_2026_update():
    """Structured facts transcribed from NSO's official Q3/9M 2026 release; estimates stay labeled."""
    source = "https://www.nso.gov.vn/du-lieu-va-so-lieu-thong-ke/2026/10/thong-cao-bao-chi-tinh-hinh-kinh-te-xa-hoi-quy-iii-va-9-thang-nam-2026/"
    rows = [
        # Q3 broad sectors; contribution is the share of total real value-added growth, not GDP share.
        ("2026-Q3", "Quý III/2026", "agriculture", "Nông, lâm, thủy sản", 4.21, 5.14, np.nan, "q3_broad"),
        ("2026-Q3", "Quý III/2026", "industry", "Công nghiệp và xây dựng", 12.50, 51.57, np.nan, "q3_broad"),
        ("2026-Q3", "Quý III/2026", "services", "Dịch vụ", 9.54, 43.29, np.nan, "q3_broad"),
        # Nine-month broad sectors and named subsectors from the latest release.
        ("2026-9M", "9 tháng/2026", "agriculture", "Nông, lâm, thủy sản", 4.02, 5.35, 10.65, "9m_broad"),
        ("2026-9M", "9 tháng/2026", "agriculture", "Nông nghiệp", 3.72, 3.60, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "agriculture", "Lâm nghiệp", 3.99, 0.25, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "agriculture", "Thủy sản", 4.99, 1.50, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "industry", "Công nghiệp và xây dựng", 11.21, 49.62, 38.35, "9m_broad"),
        ("2026-9M", "9 tháng/2026", "industry", "Công nghiệp", 11.02, 41.39, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "industry", "Công nghiệp chế biến, chế tạo", 11.36, 33.85, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "industry", "Xây dựng", 12.22, 8.23, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "services", "Dịch vụ", 8.69, 45.03, 42.97, "9m_broad"),
        ("2026-9M", "9 tháng/2026", "services", "Vận tải, kho bãi", 11.03, 7.21, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "services", "Bán buôn và bán lẻ", 9.85, 11.99, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "services", "Tài chính, ngân hàng và bảo hiểm", 9.45, 5.27, np.nan, "9m_subsector"),
        ("2026-9M", "9 tháng/2026", "services", "Lưu trú và ăn uống", 8.83, 2.69, np.nan, "9m_subsector"),
    ]
    result = pd.DataFrame(rows, columns=["ky", "ky_hien_thi", "khu_vuc_ma", "hoat_dong_kinh_te",
                                        "tang_truong_gdp_thuc_yoy_pct", "ty_trong_dong_gop_tang_truong_pct",
                                        "ty_trong_gdp_9_thang_pct", "cap_phan_tich"])
    result["trang_thai"] = "uoc_tinh"
    result["don_vi_tang_truong"] = "% so với cùng kỳ"
    result["dien_giai_dong_gop"] = "Tỷ trọng (%) của mức tăng tổng giá trị tăng thêm; không phải điểm phần trăm tăng GDP."
    result["nguon_url"] = source
    result["ghi_chu"] = "Số NSO ước tính. Cơ cấu GDP 9 tháng chỉ công bố cho ba khu vực; không suy ra tỷ trọng GDP cho tiểu ngành."
    save(result, "du_lieu_gdp_nganh_chi_tiet_2026.csv")
    return result


def build_employment():
    raw = read_px("employment_activity")
    records = []
    for _, row in raw.iterrows():
        activity = str(row.iloc[0])
        for col in raw.columns[1:]:
            year = year_col(col)
            value = num(row[col])
            if year is None or not np.isfinite(value):
                continue
            is_share = str(col).lower().startswith("structure")
            records.append({"nam": year, "hoat_dong_kinh_te": activity,
                            "so_lao_dong_nghin_nguoi": np.nan if is_share else value,
                            "ty_trong_lao_dong_pct": value if is_share else np.nan,
                            "don_vi": "nghìn người" if not is_share else "% tổng lao động",
                            "loai_chi_tieu": "co_cau" if is_share else "so_luong",
                            "trang_thai": status_for(col, year)})
    result = pd.DataFrame(records)
    save(result, "du_lieu_lao_dong_theo_nganh.csv")
    return result


def build_productivity():
    raw = read_px("labor_productivity")
    records = []
    for _, row in raw.iterrows():
        industry = str(row.iloc[0])
        for col in raw.columns[1:]:
            year = year_col(col)
            value = num(row[col])
            if year is None or not np.isfinite(value):
                continue
            records.append({"nam": year, "nganh": industry, "nang_suat_lao_dong_danh_nghia_trieu_vnd_nguoi": value,
                            "don_vi": "triệu VND/lao động (giá hiện hành theo bảng NSO)",
                            "trang_thai": status_for(col, year),
                            "ghi_chu": "Chỉ tiêu danh nghĩa theo mức giá của năm; không coi tăng trưởng này là tăng năng suất thực."})
    result = pd.DataFrame(records).sort_values(["nganh", "nam"])
    result["tang_truong_danh_nghia_yoy_pct"] = result.groupby("nganh")["nang_suat_lao_dong_danh_nghia_trieu_vnd_nguoi"].pct_change() * 100
    save(result, "du_lieu_nang_suat_lao_dong.csv")
    return result


def build_unemployment():
    raw = read_px("unemployment_region")
    records = []
    for _, row in raw.iterrows():
        residence, region = str(row.iloc[0]), str(row.iloc[1])
        for col in raw.columns[2:]:
            year = year_col(col)
            value = num(row[col])
            if year is not None and np.isfinite(value):
                records.append({"nam": year, "khu_vuc_cu_tru": residence, "vung_kinh_te": region,
                                "ty_le_that_nghiep_luc_luong_lao_dong_trong_do_tuoi_pct": value,
                                "trang_thai": status_for(col, year)})
    result = pd.DataFrame(records)
    save(result, "du_lieu_that_nghiep_vung.csv")
    return result


def build_investment():
    raw = read_px("investment_ownership")
    records = []
    for _, row in raw.iterrows():
        year = year_col(row["Year"])
        if year is None:
            continue
        item = str(row["Items"])
        is_share = item.lower().startswith("structure")
        is_index = item.lower().startswith("index")
        is_constant = "constant" in item.lower()
        measure = "ty_trong_pct" if is_share else ("chi_so_tang_truong_100" if is_index else ("gia_tri_thuc_2010_vnd_ty" if is_constant else "gia_tri_hien_hanh_vnd_ty"))
        for source_col, owner in [("Total", "Tổng"), ("State", "Nhà nước"), ("Non- State", "Ngoài nhà nước"), ("Foreign invested sector", "Khu vực FDI")]:
            value = num(row[source_col])
            if np.isfinite(value):
                records.append({"nam": year, "loai_gia_tri": measure, "loai_hinh_so_huu": owner,
                                "gia_tri": value, "don_vi": "%" if is_share else ("chỉ số năm trước = 100" if is_index else ("tỷ VND giá cố định 2010" if is_constant else "tỷ VND hiện hành")),
                                "trang_thai": status_for(row["Year"], year)})
    ownership = pd.DataFrame(records)
    save(ownership, "du_lieu_dau_tu_theo_loai_hinh.csv")

    raw = read_px("investment_constant_industry")
    industry_records = []
    for _, row in raw.iterrows():
        industry = str(row.iloc[0])
        for col in raw.columns[1:]:
            year = year_col(col)
            value = num(row[col])
            if year is not None and np.isfinite(value):
                industry_records.append({"nam": year, "nganh": industry, "dau_tu_thuc_vnd_ty_gia_2010": value,
                                         "don_vi": "tỷ VND giá cố định 2010", "trang_thai": status_for(col, year)})
    investment_industry = pd.DataFrame(industry_records)
    investment_industry["tang_truong_yoy_pct"] = investment_industry.groupby("nganh")["dau_tu_thuc_vnd_ty_gia_2010"].pct_change() * 100
    save(investment_industry, "du_lieu_dau_tu_theo_nganh_thuc.csv")
    return ownership, investment_industry


def build_fiscal():
    records = []
    for source, name, label_field, measure_label in [
        ("budget_revenue", "thu_ngan_sach", "State budget revenue", "thu"),
        ("budget_expenditure", "chi_ngan_sach", "State budget expenditure", "chi"),
    ]:
        raw = read_px(source)
        year_cols = [(c, year_col(c)) for c in raw.columns if year_col(c)]
        for _, row in raw.iterrows():
            measure = str(row.iloc[0])
            indicator = str(row.iloc[1])
            is_share = "structure" in measure.lower() or "co c?u" in measure.lower()
            for col, year in year_cols:
                value = num(row[col])
                if year is None or not np.isfinite(value):
                    continue
                records.append({"nam": year, "phan_he": name, "loai_thu_chi": measure_label,
                                "chi_tieu": indicator, "loai_gia_tri": "ty_trong_pct" if is_share else "gia_tri_hien_hanh_vnd_ty",
                                "gia_tri": value, "don_vi": "%" if is_share else "tỷ VND hiện hành",
                                "trang_thai": status_for(col, year)})
    result = pd.DataFrame(records)
    save(result, "du_lieu_tai_khoa_ngan_sach.csv")
    return result


def build_credit_and_rates():
    raw = read_px("credit_outstanding")
    credit = []
    for _, row in raw.iterrows():
        measure, indicator, detail = [str(row.iloc[i]) for i in range(3)]
        for col in raw.columns[3:]:
            year = year_col(col)
            value = num(row[col])
            if year is None or not np.isfinite(value):
                continue
            growth = "growth" in measure.lower()
            credit.append({"nam": year, "chi_tieu": indicator, "ky_han_ngoai_te": detail,
                           "loai_gia_tri": "tang_truong_pct" if growth else "du_no",
                           "gia_tri": value, "don_vi": "%" if growth else "nghìn tỷ VND",
                           "trang_thai": status_for(col, year)})
    credit = pd.DataFrame(credit)
    save(credit, "du_lieu_tin_dung_du_no.csv")

    raw = read_px("interest_rates")
    rates = []
    for _, row in raw.iterrows():
        category = str(row.iloc[0])
        for col in raw.columns[1:]:
            year = year_col(col)
            value = num(row[col])
            if year is not None and np.isfinite(value):
                rates.append({"nam": year, "loai_lai_suat": category, "gia_tri_pct": value,
                              "don_vi": "phần trăm/năm", "trang_thai": status_for(col, year)})
    rates = pd.DataFrame(rates)
    save(rates, "du_lieu_lai_suat_ngan_hang.csv")
    return credit, rates


def build_bop():
    raw = read_px("balance_of_payments")
    records = []
    for _, row in raw.iterrows():
        kind, indicator = str(row.iloc[0]), str(row.iloc[1])
        for col in raw.columns[2:]:
            year = year_col(col)
            value = num(row[col])
            if year is None or not np.isfinite(value):
                continue
            percent = "percent" in indicator.lower() or "%" in indicator
            records.append({"nam": year, "chi_tieu": indicator, "gia_tri": value,
                            "don_vi": "% GDP" if percent else "triệu USD",
                            "loai_gia_tri": "ty_trong_pct" if percent else "gia_tri_usd_trieu",
                            "trang_thai": status_for(col, year)})
    result = pd.DataFrame(records)
    save(result, "du_lieu_can_can_thanh_toan.csv")
    return result


def build_tfp():
    raw = read_px("tfp_contribution")
    result = raw.rename(columns={
        "Period": "giai_doan", "Capital": "dong_gop_von_pct", "Labour": "dong_gop_lao_dong_pct",
        "Total factor of productivity (TFP)": "dong_gop_tfp_pct", "Growth rate of TFP": "tang_truong_tfp_pct",
    })
    for col in result.columns[1:]:
        result[col] = pd.to_numeric(result[col], errors="coerce")
    result["ghi_chu"] = "Tỷ trọng đóng góp theo kỳ 5 năm; không phải chuỗi năm. TFP là phần đóng góp vào tăng trưởng theo bảng NSO."
    save(result, "du_lieu_dong_gop_tfp.csv")
    return result


def build_commodities():
path = ROOT / "data" / "world_bank" / "raw" / "expanded" / "world_bank_commodity_monthly.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Monthly Prices"]
    rows = list(sheet.iter_rows(min_row=5, max_row=6, values_only=True))
    header, units = rows[0], rows[1]
    selected = {
        "Crude oil, average": "Dầu thô bình quân",
        "Coal, Australian": "Than Australia",
        "Natural gas, Europe": "Khí tự nhiên châu Âu",
        "Rice, Viet Namese 5%": "Gạo Việt Nam 5% tấm",
        "Urea ": "Urê",
        "Iron ore, cfr spot": "Quặng sắt CFR",
        "Copper": "Đồng",
    }
    colmap = {label: idx for idx, label in enumerate(header) if label in selected}
    records = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        period = str(row[0] or "")
        match = re.fullmatch(r"(\d{4})M(\d{2})", period)
        if not match:
            continue
        for label, idx in colmap.items():
            value = num(row[idx])
            if np.isfinite(value):
                records.append({"thoi_gian": f"{match.group(1)}-{match.group(2)}-01",
                                "ma_hang_hoa": label, "hang_hoa": selected[label], "gia_usd": value,
                                "don_vi_nguon": units[idx], "trang_thai": "quan_sat"})
    result = pd.DataFrame(records)
    if result.empty:
        raise ValueError("No selected World Bank commodity observations could be parsed")
    result["thoi_gian"] = pd.to_datetime(result["thoi_gian"])
    result = result.sort_values(["ma_hang_hoa", "thoi_gian"])
    result["thay_doi_yoy_pct"] = result.groupby("ma_hang_hoa")["gia_usd"].pct_change(12) * 100
    result["thay_doi_mom_pct"] = result.groupby("ma_hang_hoa")["gia_usd"].pct_change() * 100
    result["thoi_gian"] = result["thoi_gian"].dt.strftime("%Y-%m-%d")
    save(result, "du_lieu_gia_hang_hoa_quoc_te.csv")
    return result


def broad_labor_productivity(sectors, labor):
    activities = labor[labor["loai_chi_tieu"].eq("so_luong")].copy()
    industry_names = {
        "Mining and quarrying", "Manufacturing", "Electricity, gas, stream and air conditioning supply",
        "Water supply, sewerage, waste management and remediation activities", "Construction",
    }
    def broad(activity):
        if activity == "Agriculture, forestry and fishing":
            return "Nông, lâm, thủy sản"
        if activity in industry_names:
            return "Công nghiệp và xây dựng"
        if activity == "TOTAL":
            return "Tổng lao động"
        return "Dịch vụ"
    activities["khu_vuc"] = activities["hoat_dong_kinh_te"].map(broad)
    aggregated = activities[activities["khu_vuc"].ne("Tổng lao động")].groupby(["nam", "khu_vuc"], as_index=False)["so_lao_dong_nghin_nguoi"].sum()
    real = sectors[sectors["ma_nganh"].isin(["nong_nghiep", "cong_nghiep_xay_dung", "dich_vu"])].copy()
    real["khu_vuc"] = real["nganh"]
    merged = real.merge(aggregated, on=["nam", "khu_vuc"], how="inner")
    merged["nang_suat_thuc_proxy_trieu_vnd_2010_nguoi"] = merged["gdp_thuc_vnd_ty_gia_2010"] / merged["so_lao_dong_nghin_nguoi"]
    merged["tang_truong_nang_suat_thuc_proxy_yoy_pct"] = merged.sort_values("nam").groupby("khu_vuc")["nang_suat_thuc_proxy_trieu_vnd_2010_nguoi"].pct_change() * 100
    merged["ghi_chu"] = "Proxy suy ra: GDP thực giá 2010 (tỷ VND) / lao động (nghìn người) = triệu VND/lao động. Nhóm ngành đã gộp; chịu ảnh hưởng khác biệt phạm vi GDP/việc làm, không đồng nhất chỉ tiêu năng suất chính thức."
    save(merged, "du_lieu_nang_suat_thuc_theo_khu_vuc.csv")
    return merged


def build_macro_panel(contrib, sector, labor, productivity_proxy, ownership, investment_industry,
                      fiscal, credit, unemployment, bop, commodities):
panel = pd.read_csv(data_path("du_lieu_chi_so_kinh_te_hang_nam.csv"), encoding="utf-8-sig")
    panel["nam"] = pd.to_numeric(panel["nam"], errors="coerce")
    cols = ["nam", "tang_truong_gdp_thuc_pct", "lam_phat_cpi_yoy_pct", "trade_to_gdp_pct", "can_can_thuong_mai_tren_gdp_pct", "von_fdi_thuc_hien_tren_gdp_pct"]
    panel = panel[cols].copy()

    panel = panel.merge(contrib[["nam", "dong_gop_von_co_dinh_diem_pct", "dong_gop_ton_kho_diem_pct", "dong_gop_tieu_dung_nha_nuoc_diem_pct", "dong_gop_tieu_dung_ho_gia_dinh_diem_pct", "dong_gop_can_can_hang_hoa_dich_vu_diem_pct", "dong_gop_chenh_lech_thong_ke_diem_pct"]], on="nam", how="left")
    sector_growth = sector[sector["ma_nganh"].isin(["nong_nghiep", "cong_nghiep_xay_dung", "dich_vu"])].pivot(index="nam", columns="ma_nganh", values="tang_truong_yoy_pct").reset_index().rename(columns={"nong_nghiep": "tang_truong_nganh_nong_nghiep_pct", "cong_nghiep_xay_dung": "tang_truong_nganh_cong_nghiep_xay_dung_pct", "dich_vu": "tang_truong_nganh_dich_vu_pct"})
    panel = panel.merge(sector_growth, on="nam", how="left")

    own = ownership[(ownership["loai_gia_tri"].eq("gia_tri_thuc_2010_vnd_ty")) & ownership["loai_hinh_so_huu"].ne("Tổng")].pivot(index="nam", columns="loai_hinh_so_huu", values="gia_tri").reset_index()
    own.columns = ["nam"] + ["dau_tu_thuc_" + {"Nhà nước": "nha_nuoc", "Ngoài nhà nước": "ngoai_nha_nuoc", "Khu vực FDI": "fdi"}.get(str(c), str(c)) + "_vnd_ty" for c in own.columns[1:]]
    panel = panel.merge(own, on="nam", how="left")
    for col in [c for c in panel.columns if c.startswith("dau_tu_thuc_")]:
        panel[col.replace("_vnd_ty", "_tang_truong_yoy_pct")] = pd.to_numeric(panel[col], errors="coerce").pct_change() * 100
    total_inv = ownership[(ownership["loai_gia_tri"].eq("gia_tri_thuc_2010_vnd_ty")) & ownership["loai_hinh_so_huu"].eq("Tổng")][["nam", "gia_tri"]].rename(columns={"gia_tri": "dau_tu_thuc_tong_vnd_ty"}).sort_values("nam")
    total_inv["tang_truong_dau_tu_thuc_tong_yoy_pct"] = total_inv["dau_tu_thuc_tong_vnd_ty"].pct_change() * 100
    total_inv = total_inv.drop(columns=["gia_tri"], errors="ignore")
    panel = panel.merge(total_inv, on="nam", how="left")

    inv_industry_total = investment_industry[investment_industry["nganh"].eq("TOTAL")][["nam", "dau_tu_thuc_vnd_ty_gia_2010", "tang_truong_yoy_pct"]].rename(columns={"dau_tu_thuc_vnd_ty_gia_2010": "dau_tu_thuc_toan_nganh_vnd_ty", "tang_truong_yoy_pct": "tang_truong_dau_tu_nganh_yoy_pct"})
    panel = panel.merge(inv_industry_total, on="nam", how="left")

current_gdp = pd.read_csv(data_path("du_lieu_gdp_theo_chi_tieu.csv"), encoding="utf-8-sig")
    current_raw = read_px("gdp_current_expenditure")
    gdp_level = []
    total_row = current_raw[(current_raw.iloc[:, 0].str.lower().str.startswith("total")) & current_raw.iloc[:, 1].eq("TOTAL")]
    if not total_row.empty:
        row = total_row.iloc[0]
        for col in current_raw.columns[2:]:
            year = year_col(col); value = num(row[col])
            if year is not None and np.isfinite(value): gdp_level.append({"nam": year, "gdp_danh_nghia_vnd_ty": value})
    gdp_level = pd.DataFrame(gdp_level)
    panel = panel.merge(gdp_level, on="nam", how="left")
    dev_spend = fiscal[(fiscal["phan_he"].eq("chi_ngan_sach")) & fiscal["loai_gia_tri"].eq("gia_tri_hien_hanh_vnd_ty") & fiscal["chi_tieu"].str.lower().str.startswith("expenditure on development investment")][["nam", "gia_tri"]].rename(columns={"gia_tri": "chi_dau_tu_phat_trien_vnd_ty"})
    total_spend = fiscal[(fiscal["phan_he"].eq("chi_ngan_sach")) & fiscal["loai_gia_tri"].eq("gia_tri_hien_hanh_vnd_ty") & fiscal["chi_tieu"].str.lower().str.startswith("total expenditure")][["nam", "gia_tri"]].rename(columns={"gia_tri": "tong_chi_ngan_sach_vnd_ty"})
    revenue = fiscal[(fiscal["phan_he"].eq("thu_ngan_sach")) & fiscal["loai_gia_tri"].eq("gia_tri_hien_hanh_vnd_ty") & fiscal["chi_tieu"].str.strip().str.lower().eq("total(*)")][["nam", "gia_tri"]].rename(columns={"gia_tri": "tong_thu_ngan_sach_vnd_ty"})
    panel = panel.merge(dev_spend, on="nam", how="left").merge(total_spend, on="nam", how="left").merge(revenue, on="nam", how="left")
    panel["dau_tu_phat_trien_chi_ngan_sach_tren_gdp_pct"] = panel["chi_dau_tu_phat_trien_vnd_ty"] / panel["gdp_danh_nghia_vnd_ty"] * 100
    panel["ty_trong_dau_tu_phat_trien_trong_chi_pct"] = panel["chi_dau_tu_phat_trien_vnd_ty"] / panel["tong_chi_ngan_sach_vnd_ty"] * 100
    panel["can_doi_ngan_sach_don_gian_vnd_ty"] = panel["tong_thu_ngan_sach_vnd_ty"] - panel["tong_chi_ngan_sach_vnd_ty"]

    cred = credit[(credit["loai_gia_tri"].eq("du_no")) & credit["chi_tieu"].str.lower().eq("outstanding") & credit["ky_han_ngoai_te"].str.lower().eq("total")][["nam", "gia_tri"]].rename(columns={"gia_tri": "du_no_tin_dung_tong_nghin_ty_vnd"})
    cred["tang_truong_du_no_tin_dung_yoy_pct"] = cred["du_no_tin_dung_tong_nghin_ty_vnd"].pct_change() * 100
    panel = panel.merge(cred, on="nam", how="left")
    panel["du_no_tin_dung_tren_gdp_pct"] = panel["du_no_tin_dung_tong_nghin_ty_vnd"] * 1000 / panel["gdp_danh_nghia_vnd_ty"] * 100

    ca = bop[bop["chi_tieu"].str.lower().eq("current account")][["nam", "gia_tri"]].rename(columns={"gia_tri": "can_can_vang_lai_trieu_usd"})
    ca_pct = bop[bop["chi_tieu"].str.lower().str.contains("current account as percent")][["nam", "gia_tri"]].rename(columns={"gia_tri": "can_can_vang_lai_tren_gdp_pct"})
    panel = panel.merge(ca, on="nam", how="left").merge(ca_pct, on="nam", how="left")

    unemployment = unemployment[(unemployment["khu_vuc_cu_tru"].eq("WHOLE COUNTRY")) & unemployment["vung_kinh_te"].eq("Total")][["nam", "ty_le_that_nghiep_luc_luong_lao_dong_trong_do_tuoi_pct"]].rename(columns={"ty_le_that_nghiep_luc_luong_lao_dong_trong_do_tuoi_pct": "that_nghiep_do_tuoi_lao_dong_pct"})
    panel = panel.merge(unemployment, on="nam", how="left")

    proxy_growth = productivity_proxy.groupby("nam").apply(lambda df: np.average(df["tang_truong_nang_suat_thuc_proxy_yoy_pct"].dropna(), weights=df.loc[df["tang_truong_nang_suat_thuc_proxy_yoy_pct"].notna(), "so_lao_dong_nghin_nguoi"]) if df["tang_truong_nang_suat_thuc_proxy_yoy_pct"].notna().any() else np.nan, include_groups=False).rename("tang_truong_nang_suat_thuc_proxy_binh_quan_pct").reset_index()
    panel = panel.merge(proxy_growth, on="nam", how="left")

    comm = commodities.copy()
    comm["nam"] = pd.to_datetime(comm["thoi_gian"]).dt.year
    annual = comm.groupby(["nam", "hang_hoa"], as_index=False)["gia_usd"].mean().rename(columns={"gia_usd": "gia_binh_quan_nam_usd"})
    annual = annual.sort_values(["hang_hoa", "nam"])
    annual["thay_doi_binh_quan_nam_yoy_pct"] = annual.groupby("hang_hoa")["gia_binh_quan_nam_usd"].pct_change() * 100
    chosen = {"Dầu thô bình quân": "thay_doi_gia_dau_binh_quan_nam_yoy_pct", "Than Australia": "thay_doi_gia_than_australia_nam_yoy_pct", "Gạo Việt Nam 5% tấm": "thay_doi_gia_gao_viet_nam_nam_yoy_pct", "Urê": "thay_doi_gia_ure_nam_yoy_pct", "Quặng sắt CFR": "thay_doi_gia_quang_sat_nam_yoy_pct", "Đồng": "thay_doi_gia_dong_nam_yoy_pct", "Khí tự nhiên châu Âu": "thay_doi_gia_khi_chau_au_nam_yoy_pct"}
    annual = annual[annual["hang_hoa"].isin(chosen)]
    commodities_panel = annual.pivot(index="nam", columns="hang_hoa", values="thay_doi_binh_quan_nam_yoy_pct").reset_index().rename(columns=chosen)
    panel = panel.merge(commodities_panel, on="nam", how="left").sort_values("nam")
    panel["trang_thai"] = panel["nam"].map(lambda y: "uoc_tinh" if y == 2025 else "theo_bang_nguon")
    save(panel, "du_lieu_so_sanh_dong_luc_gdp.csv")

    pair_specs = [
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng đầu tư thực toàn nền kinh tế", "tang_truong_dau_tu_thuc_tong_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng đầu tư thực khu vực nhà nước", "dau_tu_thuc_nha_nuoc_tang_truong_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng vốn FDI thực", "dau_tu_thuc_fdi_tang_truong_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng tín dụng danh nghĩa", "tang_truong_du_no_tin_dung_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Lạm phát CPI", "lam_phat_cpi_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Thất nghiệp trong độ tuổi", "that_nghiep_do_tuoi_lao_dong_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng năng suất thực proxy bình quân", "tang_truong_nang_suat_thuc_proxy_binh_quan_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng giá dầu bình quân năm", "thay_doi_gia_dau_binh_quan_nam_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tăng giá gạo Việt Nam 5%", "thay_doi_gia_gao_viet_nam_nam_yoy_pct"),
        ("Tăng GDP thực", "tang_truong_gdp_thuc_pct", "Tỷ trọng cán cân vãng lai/GDP", "can_can_vang_lai_tren_gdp_pct"),
    ]
    corr = []
    for xname, xcol, yname, ycol in pair_specs:
        pair = panel[["nam", xcol, ycol]].dropna()
        if len(pair) < 3:
            pearson = spearman = np.nan
        else:
            pearson = pair[xcol].corr(pair[ycol], method="pearson")
            spearman = pair[xcol].corr(pair[ycol], method="spearman")
        corr.append({"bien_x": xname, "truong_x": xcol, "bien_y": yname, "truong_y": ycol,
                     "pearson": pearson, "spearman": spearman, "so_quan_sat": len(pair),
                     "giai_doan": f"{int(pair.nam.min())}-{int(pair.nam.max())}" if len(pair) else "",
                     "dien_giai": "Tương quan mô tả trên biến tốc độ/tỷ lệ; không chứng minh nhân quả; mẫu ngắn và có thể gồm số sơ bộ/ước tính."})
    corr = pd.DataFrame(corr)
    save(corr, "du_lieu_tuong_quan_dong_luc_gdp.csv")
    return panel, corr


def quality_report(outputs, employment, contrib, productivity_proxy, tfp, detailed_activities):
    checks = []
    def add(name, ok, details, rows=None):
        checks.append({"kiem_tra": name, "trang_thai": "dat" if ok else "can_xem_lai", "so_dong": rows if rows is not None else "", "ghi_chu": details})
    for name, df in outputs.items():
        key_cols = [c for c in ["nam", "thoi_gian", "ky", "chi_tieu", "loai_chi_tieu", "loai_gia_tri", "phan_he", "ky_han_ngoai_te", "nganh", "hoat_dong_kinh_te", "hoat_dong_kinh_te_nguon", "khu_vuc_ma", "cap_phan_tich", "loai_lai_suat", "bien_x", "bien_y", "khu_vuc_cu_tru", "vung_kinh_te", "loai_hinh_so_huu", "hang_hoa", "giai_doan"] if c in df.columns]
        dup = int(df.duplicated(key_cols).sum()) if key_cols else 0
        add(f"{name}: khóa trùng", dup == 0, f"Khóa: {', '.join(key_cols)}; trùng={dup}.", len(df))
        add(f"{name}: bảng không rỗng", len(df) > 0, f"Số dòng chuẩn hóa: {len(df)}.", len(df))
    shares = employment[employment["loai_chi_tieu"].eq("co_cau") & employment["hoat_dong_kinh_te"].ne("TOTAL")].groupby("nam")["ty_trong_lao_dong_pct"].sum()
    mismatch = shares[(shares - 100).abs() > 0.25]
    add("Cơ cấu lao động cộng xấp xỉ 100%", mismatch.empty, "Tổng tỷ trọng theo ngành mỗi năm; sai lệch cho phép 0,25 điểm % do làm tròn. " + (f"Lệch: {mismatch.to_dict()}" if not mismatch.empty else ""), len(shares))
    add("Phân rã GDP theo chi tiêu", (contrib["sai_lech_phan_ra_diem_pct"].dropna().abs() < 0.15).all(), "Tổng đóng góp các thành phần rời rạc so với tốc độ tăng GDP chính thức; sai lệch trong 0,15 điểm % cho phép làm tròn.", len(contrib))
    add("Năng suất proxy có quan sát dương", (productivity_proxy["nang_suat_thuc_proxy_trieu_vnd_2010_nguoi"] > 0).all(), "GDP thực chia số lao động đã gộp theo ba khu vực; proxy mô tả, không thay chỉ tiêu chính thức.", len(productivity_proxy))
    share_sum = tfp[["dong_gop_von_pct", "dong_gop_lao_dong_pct", "dong_gop_tfp_pct"]].sum(axis=1)
    add("Đóng góp TFP + vốn + lao động xấp xỉ 100%", ((share_sum - 100).abs() < 0.15).all(), "Kiểm tra phân rã theo ba kỳ 5 năm; sai số do làm tròn.", len(tfp))
    current_source = read_px("gdp_current_activity")
    official_totals = {}
    for _, row in current_source[current_source.iloc[:, 1].astype(str).str.strip().eq("TOTAL") & current_source.iloc[:, 0].astype(str).str.startswith("Value")].iterrows():
        for col in current_source.columns[2:]:
            year = year_col(col); value = num(row[col])
            if year is not None and np.isfinite(value): official_totals[year] = value
    observed_totals = detailed_activities.groupby("nam")["gdp_hien_hanh_ty_vnd"].sum()
    common = sorted(set(observed_totals.index) & set(official_totals))
    level_error = max((abs(observed_totals[y] - official_totals[y]) / official_totals[y] * 100 for y in common), default=np.inf)
    add("GDP danh nghĩa theo hoạt động cộng bằng tổng NSO", level_error < 0.05,
        f"Chênh tương đối lớn nhất giữa tổng ngành + thuế và GDP tổng NSO: {level_error:.5f}% (ngưỡng 0,05%).", len(common))
    share_sums = detailed_activities.groupby("nam")["ty_trong_gdp_hien_hanh_pct"].sum().dropna()
    share_error = float((share_sums - 100).abs().max()) if not share_sums.empty else np.inf
    add("Cơ cấu GDP chi tiết gồm thuế cộng xấp xỉ 100%", share_error < 0.15,
        f"Sai lệch tuyệt đối lớn nhất của tỷ trọng ngành + thuế so với 100%: {share_error:.5f} điểm % (ngưỡng 0,15 điểm %).", len(share_sums))
    stable_contrib = detailed_activities[detailed_activities["nam"].ne(2010)].groupby("nam")["dong_gop_gdp_tang_truong_diem_pct"].sum(min_count=1)
    real_source = read_px("gdp_constant_growth_activity")
    official_growth = {}
    for _, row in real_source[row_filter := (real_source.iloc[:, 1].astype(str).str.strip().eq("TOTAL")) & real_source.iloc[:, 0].astype(str).str.startswith("Index")].iterrows():
        for col in real_source.columns[2:]:
            year = year_col(col); value = num(row[col])
            if year is not None and np.isfinite(value): official_growth[year] = value - 100
    growth_years = [y for y in stable_contrib.index if y in official_growth and np.isfinite(stable_contrib[y])]
    contribution_error = max((abs(stable_contrib[y] - official_growth[y]) for y in growth_years), default=np.inf)
    add("Đóng góp GDP thực chi tiết cộng khớp tăng trưởng NSO", contribution_error < 0.20,
        f"Sai lệch lớn nhất giữa tổng đóng góp theo giá cố định và chỉ số tăng GDP NSO: {contribution_error:.4f} điểm % (ngưỡng 0,20; bỏ năm đổi mốc 2010).", len(growth_years))
    result = pd.DataFrame(checks)
    save(result, "kiem_tra_du_lieu_mo_rong.csv")
    return result


def build_analysis_report(panel, corr, contrib, sectors, proxy, ownership, fiscal, credit, bop, commodities, tfp, quality):
    p = panel.set_index("nam")
    latest = int(panel["nam"].max())
    def v(col, year=latest):
        return p.loc[year, col] if year in p.index and col in p.columns else np.nan
    def fmt(value, digits=2):
        return "không có" if not np.isfinite(value) else f"{value:,.{digits}f}"
    latest_contrib = contrib[contrib["nam"].eq(latest)].iloc[0] if (contrib["nam"] == latest).any() else None
    sectors_latest = sectors[sectors["nam"].eq(latest)].set_index("nganh")
    agri_first = proxy[(proxy["nam"].eq(proxy["nam"].min())) & proxy["khu_vuc"].eq("Nông, lâm, thủy sản")]
    agri_last = proxy[(proxy["nam"].eq(proxy["nam"].max())) & proxy["khu_vuc"].eq("Nông, lâm, thủy sản")]
    if not agri_first.empty and not agri_last.empty:
        agri_prod_change = agri_last.iloc[0]["nang_suat_thuc_proxy_trieu_vnd_2010_nguoi"] / agri_first.iloc[0]["nang_suat_thuc_proxy_trieu_vnd_2010_nguoi"] - 1
        productivity_span = f"{int(agri_first.iloc[0]['nam'])}–{int(agri_last.iloc[0]['nam'])}"
    else:
        agri_prod_change = np.nan; productivity_span = ""
    latest_ownership = ownership[(ownership["nam"].eq(latest)) & ownership["loai_gia_tri"].eq("ty_trong_pct") & ownership["loai_hinh_so_huu"].ne("Tổng")]
    latest_ownership_text = ", ".join(f"{r.loai_hinh_so_huu}: {r.gia_tri:.1f}%" for r in latest_ownership.itertuples(index=False)) or "không có"
    latest_ca = bop[bop["chi_tieu"].str.lower().eq("current account as percent of gdp (%)")]
    latest_ca_val = latest_ca[latest_ca["nam"].eq(latest)]["gia_tri"].iloc[0] if (latest_ca["nam"] == latest).any() else np.nan
    commodity_dates = pd.to_datetime(commodities["thoi_gian"])
    commodity_latest = commodity_dates.max().strftime("%Y-%m")
    commodity_names = commodities[commodities["thoi_gian"].eq(commodity_dates.max().strftime("%Y-%m-%d"))]["hang_hoa"].nunique()
    corr_top = corr.dropna(subset=["pearson"]).assign(abs_r=lambda x: x["pearson"].abs()).sort_values("abs_r", ascending=False).head(4)
    corr_lines = "\n".join(f"- {r.bien_x} × {r.bien_y}: Pearson {r.pearson:.2f}, Spearman {r.spearman:.2f}, n={int(r.so_quan_sat)} ({r.giai_doan}). Đây là tương quan mô tả, không phải tác động nhân quả." for r in corr_top.itertuples(index=False))
    tfp_lines = "\n".join(f"- {r.giai_doan}: TFP đóng góp {r.dong_gop_tfp_pct:.2f}%, vốn {r.dong_gop_von_pct:.2f}%, lao động {r.dong_gop_lao_dong_pct:.2f}%; tăng trưởng TFP {r.tang_truong_tfp_pct:.2f}%." for r in tfp.itertuples(index=False))
    latest_leafs = []
    if latest_contrib is not None:
        for label, key in [("Vốn cố định", "von_co_dinh"), ("Tồn kho", "ton_kho"), ("Tiêu dùng nhà nước", "tieu_dung_nha_nuoc"), ("Tiêu dùng hộ gia đình", "tieu_dung_ho_gia_dinh"), ("Cán cân hàng hóa & dịch vụ", "can_can_hang_hoa_dich_vu"), ("Sai số thống kê", "chenh_lech_thong_ke")]:
            value = latest_contrib.get(f"dong_gop_{key}_diem_pct", np.nan)
            if np.isfinite(value): latest_leafs.append(f"- {label}: {value:+.2f} điểm phần trăm")
    contribution_text = "\n".join(latest_leafs) if latest_leafs else "Không đủ dữ liệu"
    report = f"""# Phân tích dữ liệu mở rộng: động lực GDP Việt Nam

## Tóm tắt kỳ mới nhất

- GDP thực theo bảng NSO (giá cố định 2010) năm {latest}: tăng **{fmt(v('tang_truong_gdp_thuc_pct'))}%**; kỳ 2025 trong bảng NSO mang trạng thái ước tính. Chuỗi mục tiêu World Bank trong hệ thống vẫn giữ riêng, không thay thế nhau do khác nguồn, đơn vị và phương pháp.
- Tăng trưởng theo phương pháp chi tiêu, đóng góp vào tăng GDP {latest}:
{contribution_text}
- Cơ cấu tăng GDP sản xuất năm {latest}: nông nghiệp {fmt(sectors_latest.loc['Nông, lâm, thủy sản','tang_truong_yoy_pct'])}%; công nghiệp-xây dựng {fmt(sectors_latest.loc['Công nghiệp và xây dựng','tang_truong_yoy_pct'])}%; dịch vụ {fmt(sectors_latest.loc['Dịch vụ','tang_truong_yoy_pct'])}%. Đây là các tốc độ tăng của giá trị gia tăng thực, không phải mức đóng góp điểm phần trăm.
- Cán cân vãng lai/GDP năm {latest}: {fmt(latest_ca_val)}% theo bảng cán cân thanh toán NSO.
- Cơ cấu đầu tư theo sở hữu (giá hiện hành, kỳ mới nhất trong bảng): {latest_ownership_text}. Tỷ trọng dùng tổng đầu tư làm mẫu số; có trạng thái sơ bộ/ước tính theo nguồn.
- Dữ liệu giá hàng hóa quốc tế: tới {commodity_latest}, gồm {commodity_names} chuỗi được chọn (dầu, than, khí, gạo, urê, quặng sắt, đồng); đơn vị USD danh nghĩa từ Pink Sheet.

## Những lớp phân tích được bổ sung

1. **Chi tiêu GDP thực:** GDP, tiêu dùng hộ gia đình/nhà nước, tích lũy tài sản cố định, thay đổi tồn kho, cán cân hàng hóa-dịch vụ và sai số thống kê; tính đóng góp điểm phần trăm bằng `100 × (X_t − X_(t−1)) / GDP_(t−1)`. Các thành phần lá được dùng để tránh cộng trùng tiêu dùng/tổng đầu tư với tiểu mục.
2. **GDP theo sản xuất:** mức giá cố định 2010 và chỉ số tăng trưởng năm trước theo ngành; tách mức tăng của ngành khỏi tỷ trọng GDP danh nghĩa.
3. **Lao động và năng suất:** cơ cấu/số lao động theo ngành, thất nghiệp theo vùng và chỉ tiêu năng suất danh nghĩa của NSO. Bổ sung proxy năng suất thực ba khu vực bằng GDP thực ngành chia việc làm; proxy này là chỉ báo xấp xỉ, không thay chuỗi năng suất chính thức.
4. **Đầu tư:** vốn giải ngân theo sở hữu ở cả giá hiện hành, giá cố định 2010, tỷ trọng và tốc độ tăng; vốn thực theo ngành dùng để so sánh phân bổ vốn.
5. **Tài khóa và tiền tệ:** thu/chi ngân sách, chi đầu tư phát triển, tín dụng theo kỳ hạn/loại tiền, lãi suất tiền gửi/cho vay và tỷ lệ tín dụng/GDP danh nghĩa.
6. **Đối ngoại và cú sốc đầu vào:** cán cân thanh toán, tài khoản vãng lai, các mức giá hàng hóa thế giới và biến động theo năm.
7. **TFP:** phân rã NSO theo kỳ 5 năm, không nội suy thành chuỗi năm.

## So sánh mô tả với tăng trưởng GDP thực

{corr_lines}

Các hệ số được tính trên cặp quan sát đầy đủ từng cặp; n nhỏ và các biến vĩ mô có thể đồng thời chịu cú sốc chung, xu hướng, thay đổi phương pháp đo hoặc trạng thái sơ bộ. Không dùng các tương quan này làm bằng chứng nhân quả hoặc khẳng định một biến giúp dự báo ngoài mẫu.

## Diễn biến TFP theo kỳ NSO

{tfp_lines}

## Proxy năng suất thực theo khu vực

Từ dữ liệu việc làm và GDP thực cùng kỳ, proxy năng suất nông nghiệp tăng khoảng **{fmt(agri_prod_change*100)}%** trong giai đoạn {productivity_span}. Đây là phép chia hai chuỗi gộp, chịu ảnh hưởng thay đổi phạm vi ngành và ước tính việc làm; không nên so sánh trực tiếp với năng suất danh nghĩa của NSO.

## Giới hạn sử dụng và độ tin cậy

- Các bảng NSO ghi rõ **ước tính/sơ bộ** ở năm 2025; nhãn này được giữ trong các CSV chuẩn hóa và cần cập nhật khi NSO hiệu chỉnh.
- Bảng đầu tư, GDP và tài khóa có thể khác năm gốc, phạm vi thống kê và khái niệm giải ngân/chi ngân sách. Không ghép giá hiện hành với chuỗi thực để tính tăng trưởng thực.
- Năng suất ngành NSO được lưu riêng với nhãn giá hiện hành; chỉ proxy mới dùng GDP giá cố định 2010.
- GDP năm là quan sát năm. Không tạo số GDP tháng/quý nội suy và không đưa biến giải thích chưa được backtest vào môודל dự báo GDP hiện có.
- Hệ thống dự báo đến 2030 giữ nguyên mô hình/rolling-origin hiện hành; các bộ mới mở rộng EDA và so sánh. Muốn đưa biến ngoại sinh vào mô hình cần dữ liệu phát hành theo thời gian thực (vintage), thiết kế dự báo đúng thời điểm, kiểm định rolling-origin, so sánh với baseline và kiểm soát rò rỉ dữ liệu.
- Giá hàng hóa World Bank là giá USD danh nghĩa quốc tế, không phải giá nhập khẩu thực tế của Việt Nam hay giá bán trong nước.

## Nguồn

- [Danh mục Tài khoản quốc gia NSO]({NSO_ACCOUNTS}) — GDP theo sản xuất/chi tiêu, tài khóa, tín dụng, lãi suất, cán cân thanh toán, TFP.
- [Danh mục Lao động NSO]({NSO_LABOR}) — việc làm, năng suất lao động, thất nghiệp.
- [Danh mục Đầu tư NSO]({NSO_INVESTMENT}) — giải ngân theo sở hữu/ngành.
- [World Bank Commodity Markets / Pink Sheet]({WORLD_BANK_COMMODITIES}) — giá hàng hóa theo tháng.

## QA

- {int((quality['trang_thai']=='dat').sum())}/{len(quality)} kiểm tra mở rộng đạt; các điều kiện kiểm tra gồm khóa trùng, bảng không rỗng, tổng tỷ trọng lao động, phân rã tăng trưởng, proxy dương và tổng TFP + vốn + lao động.
- Dữ liệu thô PXWeb ở `data/nso/raw/expanded/`, Pink Sheet ở `data/world_bank/raw/expanded/`; manifest ghi URL và thời điểm tải.
"""
    (ROOT / "phan_tich_du_lieu_mo_rong.md").write_text(report, encoding="utf-8")


def build_detailed_gdp_analysis_report(activities, updates):
    latest_year = int(activities["nam"].max())
    latest = activities[activities["nam"].eq(latest_year)]
    sector_keys = [("agriculture", "Nông, lâm, thủy sản"), ("industry", "Công nghiệp và xây dựng"), ("services", "Dịch vụ")]
    lines = []
    for key, label in sector_keys:
        rows = latest[latest["khu_vuc_ma"].eq(key)].sort_values("ty_trong_gdp_hien_hanh_pct", ascending=False)
        amount = rows["gdp_hien_hanh_ty_vnd"].sum()
        share = rows["ty_trong_gdp_hien_hanh_pct"].sum()
        lines.extend([f"## {label}", "", f"GDP hiện hành năm {latest_year}: **{amount/1000:,.2f} nghìn tỷ VND**, chiếm **{share:.2f}% GDP**.", "",
                      "| Hoạt động | Tỷ trọng GDP hiện hành | Tăng GDP thực năm gần nhất | Đóng góp GDP thực vào tăng trưởng chung | Thay đổi tỷ trọng 10 năm |",
                      "|---|---:|---:|---:|---:|"])
        for row in rows.itertuples(index=False):
            real_growth = "—" if not np.isfinite(row.tang_truong_gdp_thuc_yoy_pct) else f"{row.tang_truong_gdp_thuc_yoy_pct:.2f}%"
            contribution = "—" if not np.isfinite(row.dong_gop_gdp_tang_truong_diem_pct) else f"{row.dong_gop_gdp_tang_truong_diem_pct:.2f} điểm %"
            share_change = "—" if not np.isfinite(row.bien_doi_ty_trong_10_nam_diem_pct) else f"{row.bien_doi_ty_trong_10_nam_diem_pct:+.2f} điểm %"
            lines.append(f"| {row.hoat_dong_kinh_te} | {row.ty_trong_gdp_hien_hanh_pct:.2f}% | {real_growth} | {contribution} | {share_change} |")
        lines.append("")

    recent = updates[(updates["ky"].eq("2026-9M")) & updates["cap_phan_tich"].eq("9m_subsector")]
    lines.extend(["## Các hoạt động được NSO nêu riêng trong 9 tháng 2026", "",
                  "Tốc độ là GDP thực cùng kỳ. Cột đóng góp là tỷ trọng của hoạt động trong mức tăng tổng giá trị tăng thêm; không phải điểm phần trăm tăng GDP.", "",
                  "| Khu vực | Hoạt động | Tăng GDP thực | Tỷ trọng đóng góp |", "|---|---|---:|---:|"])
    for row in recent.itertuples(index=False):
        label = dict(sector_keys).get(row.khu_vuc_ma, row.khu_vuc_ma)
        lines.append(f"| {label} | {row.hoat_dong_kinh_te} | {row.tang_truong_gdp_thuc_yoy_pct:.2f}% | {row.ty_trong_dong_gop_tang_truong_pct:.2f}% |")

    vol = activities[(activities["nam"].between(latest_year-9, latest_year)) & activities["khu_vuc_ma"].isin(["agriculture", "industry", "services"])].groupby(["khu_vuc_ma", "hoat_dong_kinh_te"], as_index=False).agg(
        so_nam_tang_truong=("tang_truong_gdp_thuc_yoy_pct", "count"), bien_dong_tang_truong_pct=("tang_truong_gdp_thuc_yoy_pct", "std"),
        tang_truong_trung_binh_pct=("tang_truong_gdp_thuc_yoy_pct", "mean"))
    high_vol = vol.sort_values("bien_dong_tang_truong_pct", ascending=False).head(5)
    lines.extend(["", "## Biến động tăng trưởng theo hoạt động, 2016–2025", "",
                  "Độ biến động là độ lệch chuẩn mẫu của tốc độ tăng GDP thực năm; chỉ tính các năm có chỉ số NSO đầy đủ.", "",
                  "| Khu vực | Hoạt động | Số năm | Tăng trưởng thực bình quân | Độ lệch chuẩn |", "|---|---|---:|---:|---:|"])
    for row in high_vol.itertuples(index=False):
        label = dict(sector_keys).get(row.khu_vuc_ma, row.khu_vuc_ma)
        lines.append(f"| {label} | {row.hoat_dong_kinh_te} | {int(row.so_nam_tang_truong)} | {row.tang_truong_trung_binh_pct:.2f}% | {row.bien_dong_tang_truong_pct:.2f} điểm % |")

    lines.extend(["", "## Phương pháp và giới hạn", "",
                  "- GDP theo hoạt động giá hiện hành và tỷ trọng lấy trực tiếp từ NSO PXWeb E03.04-05; chuỗi giá cố định và chỉ số tăng trưởng thực lấy từ E03.06-07.",
                  "- Tăng trưởng thực của từng hoạt động dùng **chỉ số năm trước = 100 trừ 100**; không dùng tăng trưởng của giá trị GDP danh nghĩa.",
                  "- Đóng góp điểm phần trăm tính bằng `100 × (VAᵢ,t − VAᵢ,t−1) / GDPₜ₋₁` trên mức giá cố định; bỏ năm 2010 vì có điểm đổi mốc/đứt quãng mức giá. QA kiểm tra tổng đóng góp so với tăng trưởng GDP NSO ở các năm có thể so sánh.",
                  "- GDP nông, lâm, thủy sản trong bảng năm chỉ là một hoạt động gộp. Không chia giá trị GDP năm theo ngành trồng trọt/chăn nuôi/lâm nghiệp/thủy sản khi nguồn không công bố mức giá trị tăng thêm tương ứng.",
                  "- Số liệu 2025 là ước tính NSO. Số 2026 mới nhất là tăng trưởng lũy kế 9 tháng và Q3, không phải GDP năm đầy đủ; tỷ trọng đóng góp không phải tỷ trọng quy mô GDP.",
                  "- IIP có ngành chi tiết hơn nhưng đo chỉ số sản lượng công nghiệp, không phải GDP/giá trị tăng thêm; dùng để tham khảo chéo chứ không gộp vào GDP.",
                  "", "## Nguồn", "",
                  "- [NSO PXWeb E03.04-05 — GDP giá hiện hành theo hoạt động kinh tế](https://pxweb.nso.gov.vn/pxweb/en/E%2003.%20National%20Accounts%20and%20State%20budget/E%2003.%20National%20Accounts%20and%20State%20budget__Px_E_03/E03.04-05.px/)",
                  "- [NSO PXWeb E03.06-07 — GDP giá cố định theo hoạt động và chỉ số năm trước = 100](https://pxweb.nso.gov.vn/pxweb/en/E%2003.%20National%20Accounts%20and%20State%20budget/E%2003.%20National%20Accounts%20and%20State%20budget__Px_E_03/E03.06-07.px/)",
                  "- [Thông cáo NSO quý III và 9 tháng 2026](https://www.nso.gov.vn/du-lieu-va-so-lieu-thong-ke/2026/10/thong-cao-bao-chi-tinh-hinh-kinh-te-xa-hoi-quy-iii-va-9-thang-nam-2026/)", ""])
    (ROOT / "phan_tich_gdp_chi_tiet_hoat_dong.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    exp, contrib = build_gdp_expenditure()
    sectors = build_real_sectors()
    detailed_activities = build_detailed_gdp_activities()
    activity_updates = build_gdp_activity_2026_update()
    employment = build_employment()
    productivity = build_productivity()
    unemployment = build_unemployment()
    ownership, investment_industry = build_investment()
    fiscal = build_fiscal()
    credit, rates = build_credit_and_rates()
    bop = build_bop()
    tfp = build_tfp()
    commodities = build_commodities()
    proxy = broad_labor_productivity(sectors, employment)
    panel, corr = build_macro_panel(contrib, sectors, employment, proxy, ownership, investment_industry, fiscal, credit, unemployment, bop, commodities)
    outputs = {
        "gdp_expenditure": exp, "gdp_contribution": contrib, "gdp_real_sectors": sectors,
        "gdp_detailed_activities": detailed_activities, "gdp_activity_2026_update": activity_updates,
        "employment": employment, "labor_productivity": productivity, "unemployment": unemployment,
        "investment_ownership": ownership, "investment_industry": investment_industry,
        "fiscal": fiscal, "credit": credit, "interest_rates": rates, "bop": bop,
        "tfp": tfp, "commodities": commodities, "productivity_proxy": proxy, "macro_panel": panel,
        "macro_correlations": corr,
    }
    quality = quality_report(outputs, employment, contrib, proxy, tfp, detailed_activities)
    build_analysis_report(panel, corr, contrib, sectors, proxy, ownership, fiscal, credit, bop, commodities, tfp, quality)
    build_detailed_gdp_analysis_report(detailed_activities, activity_updates)

    # Append extended checks to the existing QA report after the original pipeline has run.
old_path = data_path("kiem_tra_chat_luong_du_lieu.csv")
    old = pd.read_csv(old_path, encoding="utf-8-sig") if old_path.exists() else pd.DataFrame()
    extended = quality.rename(columns={"kiem_tra": "tep", "so_dong": "so_dong_du_lieu", "ghi_chu": "ghi_chu"})
    extended["so_dong_trung_khoa"] = 0
    extended["so_o_thieu_tren_khoa"] = 0
    extended["so_o_thieu"] = 0
    extended["ket_qua"] = extended["trang_thai"].map({"dat": "dat", "can_xem_lai": "can_xem_lai"})
    extended = extended[["tep", "so_dong_du_lieu", "so_dong_trung_khoa", "so_o_thieu_tren_khoa", "so_o_thieu", "ket_qua", "ghi_chu"]]
    old = old[~old.get("tep", pd.Series(dtype=str)).astype(str).str.startswith("extended::")]
    extended["tep"] = "extended::" + extended["tep"].astype(str)
    pd.concat([old, extended], ignore_index=True).to_csv(old_path, index=False, encoding="utf-8-sig", lineterminator="\n")

    print(f"Expanded data: {len(outputs)} standardized files; macro panel {len(panel)} years; {len(corr)} descriptive comparisons")
    print(f"Quality: {(quality['trang_thai']=='dat').sum()}/{len(quality)} checks passed; latest month {commodities['thoi_gian'].max()[:7]}")
    print(f"GDP expenditure decomposition max abs discrepancy: {contrib['sai_lech_phan_ra_diem_pct'].abs().max():.4f} pp")


if __name__ == "__main__":
    main()
